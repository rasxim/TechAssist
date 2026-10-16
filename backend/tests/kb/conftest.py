"""Shared fixtures: a small synthetic workbook that mimics the real cheat sheet's structures,
so tests never need the internal (gitignored) workbook."""

from pathlib import Path

import openpyxl
import pytest

SHEET_ID = "SHEETID123"


def build_workbook(path: Path, *, drop_row: str | None = None, answer_override: str | None = None) -> Path:
    wb = openpyxl.Workbook()
    cq = wb.active
    cq.title = "Common Questions"

    rows = [
        # (question, answer, answer_link)
        ("Create a Tech Lounge Ticket", "https://lafayettec.az1.qualtrics.com/jfe/form/X", "https://lafayettec.az1.qualtrics.com/jfe/form/X"),
        ("Topic", "Help URL", None),
        ("How do I connect to eduroam?", "How To Connect - Technology Help", "https://help.lafayette.edu/how-to-connect-to-the-network/"),
        ("How do I release my print job?", answer_override or "1. Upload your job\n2. Click Web Print\n3. Release it", None),
        ("How do I reset Duo?", "DUO Help Guide - Tech Lounge", "https://docs.google.com/document/d/PRIVATEDOC/edit"),
        ("Which software can I download?", f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?gid=111#gid=111", "#'Software'!A1"),
        ("How do I boot to the BIOS?", f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?gid=222#gid=222", f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?gid=222"),
        ("Where can I check warranty?", f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?gid=333#gid=333", f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?gid=333"),
        ("How do I erase a computer? (NOTE: only with approval)\n\nExtra line", None, None),
    ]
    r = 0
    for question, answer, link in rows:
        if drop_row and question == drop_row:
            continue
        r += 1
        cq.cell(r, 1, question)
        cq.cell(r, 2, answer)
        if link:
            cq.cell(r, 2).hyperlink = link if not link.startswith("#") else None
            if link.startswith("#"):
                cq.cell(r, 2).hyperlink = openpyxl.worksheet.hyperlink.Hyperlink(ref=f"B{r}", location=link[1:])

    software = wb.create_sheet("Software")
    software.cell(1, 1, "Back to Main Page").hyperlink = openpyxl.worksheet.hyperlink.Hyperlink(ref="A1", location="'Common Questions'!A1")
    software.merge_cells("A1:C1")
    software.append(["Software Title", "Apps", "URL"])
    software.append(["MATLAB", "--", "https://help.lafayette.edu/matlab/"])
    software.cell(3, 3).hyperlink = "https://help.lafayette.edu/matlab/"
    software.append(["Adobe", "Acrobat | Photoshop", "https://help.lafayette.edu/adobe/"])
    software.cell(4, 3).hyperlink = "https://help.lafayette.edu/adobe/"

    bios = wb.create_sheet("BIOS")
    bios.cell(1, 1, "Back to Main Page")
    bios.merge_cells("A1:C1")
    bios.cell(2, 1, "NOTE: Do not change settings without approval")
    bios.merge_cells("A2:C2")
    bios.append(["Brand", "BIOS Key", "Boot Key"])
    bios.append(["Dell", "F2", "F12"])
    bios.append(["Apple", "Hold Power", None])
    bios.merge_cells("B5:C5")
    bios.append(["Television", "Roku", "x"])
    bios.append(["Television", "LG", "y"])

    wb.save(path)
    return path


@pytest.fixture
def workbook(tmp_path) -> Path:
    return build_workbook(tmp_path / "cheat.xlsx")


@pytest.fixture
def gid_map(tmp_path) -> Path:
    path = tmp_path / "gid_map.yaml"
    path.write_text(
        f"""spreadsheet_id: {SHEET_ID}
gids:
  "222": {{tab: BIOS, confirmed: true, evidence: test}}
  "333": {{tab: Warranty, confirmed: false, evidence: test}}
""",
        encoding="utf-8",
    )
    return path
