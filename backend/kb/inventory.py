"""Stage 1: read the cheat sheet workbook into knowledge base records.

Pure functions, no database. Everything here is deterministic: text is copied verbatim,
and anything ambiguous gets a review flag instead of a guess.
"""

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import yaml

from .links import LinkInfo, classify, hyperlink_target

INDEX_SHEET = "Common Questions"
EMPTY_MARKERS = {"--", "-", "—"}

_BARE_URL_RE = re.compile(r"^https?://\S+$")
_WARNING_RE = re.compile(r"\bNOTE\b|\bshould NOT\b|\bdo not\b|\bdoes NOT\b|\bonly be\b|\bapproval\b", re.IGNORECASE)


@dataclass
class Source:
    source_key: str
    source_type: str  # workbook_cell | workbook_tab | web_page | private_doc
    access: str  # public | restricted | local
    title: str | None = None
    url: str | None = None
    sheet_name: str | None = None
    cell_ref: str | None = None
    snapshot_text: str | None = None

    @property
    def content_hash(self) -> str | None:
        return sha256(self.snapshot_text) if self.snapshot_text is not None else None


@dataclass
class Issue:
    slug: str
    question: str
    inline_notes: str | None
    source_key: str
    source_cell: str
    sheet_row: int
    issue_kind: str | None = None
    review_flags: list[str] = field(default_factory=list)


@dataclass
class Procedure:
    procedure_key: str
    kind: str  # steps | reference_table | policy | referral
    title: str
    content: str
    source_key: str
    source_section: str | None = None
    warnings: list[str] = field(default_factory=list)
    review_flags: list[str] = field(default_factory=list)


@dataclass
class ReferenceRow:
    procedure_key: str
    row_key: str
    position: int
    fields: dict
    source_cell: str


@dataclass
class Link:
    owner_type: str  # issue | procedure
    owner_key: str  # issue slug or procedure key
    cell_ref: str
    display_text: str | None
    url: str
    link_kind: str
    followed: bool


@dataclass
class IssueProcedure:
    issue_slug: str
    procedure_key: str
    position: int
    role: str = "primary"


@dataclass
class Inventory:
    workbook_name: str
    workbook_hash: str
    sources: dict[str, Source] = field(default_factory=dict)
    issues: list[Issue] = field(default_factory=list)
    procedures: list[Procedure] = field(default_factory=list)
    reference_rows: list[ReferenceRow] = field(default_factory=list)
    links: list[Link] = field(default_factory=list)
    issue_procedures: list[IssueProcedure] = field(default_factory=list)
    followed_tabs: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def add_source(self, source: Source) -> str:
        self.sources.setdefault(source.source_key, source)
        return source.source_key

    def counts(self) -> dict:
        return {
            "sources": len(self.sources),
            "issues": len(self.issues),
            "procedures": len(self.procedures),
            "reference_rows": len(self.reference_rows),
            "links": len(self.links),
            "issue_procedures": len(self.issue_procedures),
            "followed_tabs": len(self.followed_tabs),
        }


def sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def cell_text(value) -> str | None:
    """Cell value as verbatim text; None for empty cells and placeholder dashes."""
    if value is None:
        return None
    text = str(value)
    if not text.strip() or text.strip() in EMPTY_MARKERS:
        return None
    return text


def make_slug(question: str) -> str:
    first_line = question.strip().splitlines()[0]
    base = re.sub(r"[^a-z0-9]+", "-", first_line.lower()).strip("-")[:60].rstrip("-")
    return f"{base}-{sha256(' '.join(question.split()))[:8]}"


def find_warnings(text: str) -> list[str]:
    """Sentences or lines containing warning/authorization language, verbatim."""
    pieces = re.split(r"\n+|(?<=[.!?])\s+", text)
    return [p.strip() for p in pieces if p.strip() and _WARNING_RE.search(p)]


def is_inline_answer(text: str | None, link: LinkInfo | None) -> bool:
    """True when column B holds written instructions rather than just a link title or bare URL."""
    if not text or _BARE_URL_RE.match(text.strip()):
        return False
    if link is None:
        return True
    return "\n" in text.strip() or len(text) > 120


def load_gid_map(path: Path) -> tuple[str, dict]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    return data["spreadsheet_id"], data.get("gids", {})


def build_inventory(workbook_file: Path, gid_map_file: Path) -> Inventory:
    raw = workbook_file.read_bytes()
    inv = Inventory(workbook_name=workbook_file.name, workbook_hash=hashlib.sha256(raw).hexdigest())
    cheat_sheet_id, gid_map = load_gid_map(gid_map_file)
    wb = openpyxl.load_workbook(workbook_file, data_only=True)

    if INDEX_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook has no '{INDEX_SHEET}' sheet")

    tabs_to_follow = _read_index(inv, wb[INDEX_SHEET], cheat_sheet_id, gid_map)

    for tab in tabs_to_follow:
        if tab not in wb.sheetnames:
            inv.warnings.append(f"Linked tab '{tab}' does not exist in the workbook")
            continue
        inv.followed_tabs.append(tab)
        _read_tab(inv, wb[tab], cheat_sheet_id)

    return inv


def _read_index(inv: Inventory, ws, cheat_sheet_id: str, gid_map: dict) -> list[str]:
    """Read Common Questions. Returns the tabs its links point to, in first-seen order."""
    header_row = next(
        (r for r in range(1, ws.max_row + 1) if (cell_text(ws.cell(r, 1).value) or "").strip() == "Topic"),
        None,
    )
    if header_row is None:
        inv.warnings.append("Common Questions has no 'Topic' header row; treating every row as an issue")
        header_row = 0

    tabs: list[str] = []
    for r in range(1, ws.max_row + 1):
        if r == header_row:
            continue
        a_cell, b_cell = ws.cell(r, 1), ws.cell(r, 2)
        question = cell_text(a_cell.value)
        if question is None:
            continue

        a_ref, b_ref = f"{INDEX_SHEET}!A{r}", f"{INDEX_SHEET}!B{r}"
        source_key = inv.add_source(
            Source(f"cell:{a_ref}", "workbook_cell", "local", sheet_name=INDEX_SHEET, cell_ref=f"A{r}", snapshot_text=question)
        )
        lines = question.strip().split("\n", 1)
        issue = Issue(
            slug=make_slug(question),
            question=question,
            inline_notes=lines[1].strip() if len(lines) > 1 and lines[1].strip() else None,
            source_key=source_key,
            source_cell=a_ref,
            sheet_row=r,
            issue_kind="resource" if r < header_row else None,
        )
        if issue.inline_notes:
            issue.review_flags.append("question cell contains extra text (settings or notes); decide where it belongs")
        if find_warnings(question):
            issue.review_flags.append("question contains warning/authorization language")
        inv.issues.append(issue)

        a_target = hyperlink_target(a_cell)
        if a_target:
            info = classify(a_target, cheat_sheet_id)
            inv.links.append(Link("issue", issue.slug, a_ref, question, a_target, info.link_kind, followed=False))
            issue.review_flags.append(f"question cell has its own hyperlink ({a_target}); likely accidental")

        answer = cell_text(b_cell.value)
        b_target = hyperlink_target(b_cell)
        info = classify(b_target, cheat_sheet_id) if b_target else None
        position = 0

        if is_inline_answer(answer, info):
            proc = Procedure(
                procedure_key=f"cell:{b_ref}",
                kind="steps",
                title=lines[0].strip(),
                content=answer,
                source_key=inv.add_source(
                    Source(f"cell:{b_ref}", "workbook_cell", "local", sheet_name=INDEX_SHEET, cell_ref=f"B{r}", snapshot_text=answer)
                ),
                source_section=b_ref,
                warnings=find_warnings(answer),
            )
            if proc.warnings:
                proc.review_flags.append("warning language detected; confirm warnings")
            inv.procedures.append(proc)
            inv.issue_procedures.append(IssueProcedure(issue.slug, proc.procedure_key, position))
            position += 1

        if info is None:
            if answer is None:
                issue.review_flags.append("no answer in column B")
            continue

        followed = _resolve_index_link(inv, issue, info, answer, gid_map, tabs, position)
        inv.links.append(Link("issue", issue.slug, b_ref, answer, info.url, info.link_kind, followed))

    return tabs


def _resolve_index_link(inv, issue, info: LinkInfo, answer, gid_map, tabs, position) -> bool:
    """Decide what a Common Questions link leads to. Returns whether the pipeline follows it."""
    tab = None
    if info.link_kind == "internal_tab":
        tab = info.target_tab
    elif info.link_kind == "workbook_sheet":
        entry = gid_map.get(str(info.target_tab))
        if entry is None:
            issue.review_flags.append(f"links to an unknown tab of the online cheat sheet (gid {info.target_tab}); add it to gid_map.yaml")
            return False
        if not entry.get("confirmed"):
            issue.review_flags.append(f"gid {info.target_tab} -> '{entry['tab']}' is unconfirmed in gid_map.yaml")
            return False
        tab = entry["tab"]

    if tab:
        if tab not in tabs:
            tabs.append(tab)
        inv.issue_procedures.append(IssueProcedure(issue.slug, f"tab:{tab}", position))
        return True

    if info.link_kind == "public_doc":
        inv.add_source(Source(f"url:{info.url}", "web_page", "public", title=answer, url=info.url))
        return True
    if info.link_kind == "private_doc":
        inv.add_source(Source(f"gdoc:{info.doc_id}", "private_doc", "restricted", title=answer, url=info.url))
        issue.review_flags.append(f"procedure is in a private Google Doc; export {info.doc_id} to kb_data/private/")
        return True
    if info.link_kind == "external":
        issue.review_flags.append(f"links to an unrecognized site ({info.url}); decide whether to trust it")
    return False


def _read_tab(inv: Inventory, ws, cheat_sheet_id: str) -> None:
    """Read a linked tab as a reference table. Links inside it are recorded but never followed."""
    tab = ws.title
    proc_key = f"tab:{tab}"
    filled, full_row_merges = _merge_maps(ws)

    def value(r: int, c: int) -> str | None:
        return cell_text(filled.get((r, c), ws.cell(r, c).value))

    header: list[str] | None = None
    header_row = 0
    notes_before: list[str] = []
    notes_after: list[str] = []
    rows: list[tuple[int, list[str | None], dict]] = []
    list_mode = ws.max_column == 1

    for r in range(1, ws.max_row + 1):
        vals = [value(r, c) for c in range(1, ws.max_column + 1)]
        row_links = {}
        for c in range(1, ws.max_column + 1):
            target = hyperlink_target(ws.cell(r, c))
            if target:
                info = classify(target, cheat_sheet_id)
                inv.links.append(Link("procedure", proc_key, f"{tab}!{ws.cell(r, c).coordinate}", vals[c - 1], target, info.link_kind, followed=False))
                row_links[c] = target

        nonempty = [v for v in vals if v]
        if not nonempty:
            continue
        if r == 1 and set(nonempty) == {"Back to Main Page"}:
            continue
        is_note = not list_mode and (r in full_row_merges or (len(nonempty) == 1 and vals[0] is not None and ws.max_column > 1 and header is None))

        if list_mode:
            rows.append((r, vals, row_links))
        elif header is None and not is_note and len(nonempty) >= 2:
            header, header_row = [v or f"Column {i + 1}" for i, v in enumerate(vals)], r
        elif is_note or header is None:
            (notes_after if header else notes_before).append(nonempty[0])
        else:
            rows.append((r, vals, row_links))

    if list_mode:
        header = ["Value"]

    reference_rows = _build_reference_rows(proc_key, tab, header, rows)
    inv.reference_rows.extend(reference_rows)

    content = _render_table(header, rows, notes_before, notes_after, list_mode)
    source_key = inv.add_source(Source(proc_key, "workbook_tab", "local", title=tab, sheet_name=tab, snapshot_text=content))
    proc = Procedure(
        procedure_key=proc_key,
        kind="reference_table",
        title=tab,
        content=content,
        source_key=source_key,
        source_section=f"{tab}!header row {header_row}" if header_row else tab,
        warnings=[w for note in notes_before + notes_after for w in find_warnings(note)],
    )
    if list_mode:
        proc.review_flags.append("tab is a plain list, not a table; check the title/note lines")
    if proc.warnings:
        proc.review_flags.append("warning language detected; confirm warnings")
    if any(not link.followed and link.link_kind != "internal_tab" for link in inv.links if link.owner_key == proc_key):
        proc.review_flags.append("contains outside links that are not followed (stop rule); check the rows are usable without them")
    inv.procedures.append(proc)


def _merge_maps(ws) -> tuple[dict, set]:
    """Copy each merged range's value into every cell it covers; also return rows merged across the full width."""
    filled: dict[tuple[int, int], object] = {}
    full_rows: set[int] = set()
    for rng in ws.merged_cells.ranges:
        top_left = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                filled[(r, c)] = top_left
        if rng.min_col == 1 and rng.max_col >= ws.max_column:
            full_rows.update(range(rng.min_row, rng.max_row + 1))
    return filled, full_rows


def _build_reference_rows(proc_key: str, tab: str, header: list[str], rows) -> list[ReferenceRow]:
    def norm(v):
        return " ".join((v or "").lower().split())

    key_cols = 1
    keys = [norm(vals[0]) for _, vals, _ in rows]
    if len(set(keys)) < len(keys) and len(header) > 1:
        key_cols = 2
        keys = [f"{norm(vals[0])}|{norm(vals[1])}" for _, vals, _ in rows]

    seen: dict[str, int] = {}
    out = []
    for position, ((r, vals, row_links), key) in enumerate(zip(rows, keys)):
        if key in seen:
            key = f"{key}#row{r}"
        seen[key] = r
        fields = {header[i]: vals[i] for i in range(min(len(header), len(vals)))}
        if row_links:
            fields["_links"] = {header[c - 1]: url for c, url in row_links.items() if c - 1 < len(header)}
        out.append(ReferenceRow(proc_key, key, position, fields, f"{tab}!A{r}"))
    return out


def _render_table(header, rows, notes_before, notes_after, list_mode) -> str:
    """Markdown rendering shown to employees. Cell text is copied as-is (only pipes and line breaks escaped)."""

    def md(v):
        return (v or "").replace("|", "\\|").replace("\n", "<br>")

    parts = list(notes_before)
    if list_mode:
        parts.append("\n".join(vals[0] for _, vals, _ in rows if vals[0]))
    else:
        lines = ["| " + " | ".join(md(h) for h in header) + " |", "|" + "---|" * len(header)]
        lines += ["| " + " | ".join(md(v) for v in vals) + " |" for _, vals, _ in rows]
        parts.append("\n".join(lines))
    parts += notes_after
    return "\n\n".join(parts)
