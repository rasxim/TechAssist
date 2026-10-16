"""Checks that an import lost nothing and that access is locked down.

The workbook is recounted here with plain openpyxl, independently of inventory.py,
so a bug in the parser can't hide itself.
"""

from dataclasses import dataclass
from pathlib import Path

import openpyxl
import psycopg

from .inventory import INDEX_SHEET
from .links import hyperlink_target

KB_TABLES = [
    "kb_import_runs", "kb_sources", "kb_issues", "kb_procedures",
    "kb_issue_procedures", "kb_reference_rows", "kb_links",
]


@dataclass
class Check:
    name: str
    ok: bool
    detail: str


def run_checks(conn: psycopg.Connection, workbook_file: Path) -> list[Check]:
    wb = openpyxl.load_workbook(workbook_file, data_only=True)
    checks: list[Check] = []

    def check(name, ok, detail):
        checks.append(Check(name, bool(ok), detail))

    # 1. Every non-empty Common Questions row (except the header) is an active issue.
    ws = wb[INDEX_SHEET]
    expected_rows = {
        f"{INDEX_SHEET}!A{r}" for r in range(1, ws.max_row + 1)
        if ws.cell(r, 1).value is not None and str(ws.cell(r, 1).value).strip() not in ("", "Topic")
    }
    actual_rows = {row[0] for row in conn.execute("select source_cell from kb_issues where status <> 'inactive'")}
    missing, extra = expected_rows - actual_rows, actual_rows - expected_rows
    check("issues match Common Questions rows", not missing and not extra,
          f"{len(actual_rows)} issues for {len(expected_rows)} rows; missing={sorted(missing)[:5]} extra={sorted(extra)[:5]}")

    # 2. Every hyperlink in Common Questions and in the followed tabs is recorded, with its exact URL.
    tabs = [row[0] for row in conn.execute(
        "select sheet_name from kb_sources where source_type = 'workbook_tab' and is_active")]
    expected_links = set()
    for name in [INDEX_SHEET, *tabs]:
        for row in wb[name].iter_rows():
            for cell in row:
                target = hyperlink_target(cell)
                if target:
                    expected_links.add((f"{name}!{cell.coordinate}", target))
    actual_links = {(r[0], r[1]) for r in conn.execute("""
        select l.cell_ref, l.url from kb_links l
        left join kb_issues i on l.owner_type = 'issue' and i.id = l.owner_id
        left join kb_procedures p on l.owner_type = 'procedure' and p.id = l.owner_id
        where coalesce(i.status, p.status) <> 'inactive'""")}
    missing, extra = expected_links - actual_links, actual_links - expected_links
    check("hyperlinks preserved", not missing and not extra,
          f"{len(actual_links)} recorded for {len(expected_links)} in workbook; missing={sorted(missing)[:3]} extra={sorted(extra)[:3]}")

    # 3. Every active record traces back to an active source with a location.
    orphans = conn.execute("""
        select count(*) from kb_issues i join kb_sources s on s.id = i.source_id
        where i.status <> 'inactive' and (not s.is_active or s.cell_ref is null)""").fetchone()[0]
    check("issues have an active source cell", orphans == 0, f"{orphans} issues without one")
    orphans = conn.execute("""
        select count(*) from kb_procedures p join kb_sources s on s.id = p.source_id
        where p.status <> 'inactive' and not s.is_active""").fetchone()[0]
    check("procedures have an active source", orphans == 0, f"{orphans} procedures without one")

    # 4. Workbook procedures are verbatim copies of their source snapshot.
    mismatched = conn.execute("""
        select count(*) from kb_procedures p join kb_sources s on s.id = p.source_id
        where p.status = 'draft' and s.source_type in ('workbook_cell', 'workbook_tab')
          and p.content <> s.snapshot_text""").fetchone()[0]
    check("workbook procedures are verbatim", mismatched == 0, f"{mismatched} differ from their source text")

    # 5. Every followed tab produced rows.
    empty_tabs = [row[0] for row in conn.execute("""
        select p.title from kb_procedures p
        where p.kind = 'reference_table' and p.status <> 'inactive'
          and not exists (select 1 from kb_reference_rows r where r.procedure_id = p.id)""")]
    check("reference tabs have rows", not empty_tabs, f"empty: {empty_tabs}")

    # 6. Access control: RLS on, no policies, browser roles have no grants.
    no_rls = [row[0] for row in conn.execute(
        "select relname from pg_class where relname = any(%s) and not relrowsecurity", (KB_TABLES,))]
    check("row-level security enabled", not no_rls, f"missing on: {no_rls}")
    policies = conn.execute("select count(*) from pg_policies where tablename = any(%s)", (KB_TABLES,)).fetchone()[0]
    check("no access policies defined", policies == 0, f"{policies} policies found")
    for role in ("anon", "authenticated"):
        if conn.execute("select 1 from pg_roles where rolname = %s", (role,)).fetchone():
            readable = [t for t in KB_TABLES
                        if conn.execute("select has_table_privilege(%s, %s, 'select')", (role, t)).fetchone()[0]]
            check(f"{role} role cannot read", not readable, f"readable: {readable}")

    return checks
